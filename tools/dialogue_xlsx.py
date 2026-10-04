#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""tools/dialogue_xlsx.py —— 台詞差分表（標準格式，2026-09-29 Ray 定案）

    py tools/dialogue_xlsx.py 稿.txt -o 表.xlsx     # 由稿出一張新表
    py tools/dialogue_xlsx.py --refresh 表.xlsx       # Ray 改過差分之後：只重畫縮圖
    py tools/dialogue_xlsx.py 稿.txt --append 表.xlsx # 新的一段稿接在 Ray 那一份後面
                                                      #（他的文字一格不動，同 --refresh）

標準格式＝「羅塞爾廢城_台詞差分.xlsx」最後一版的樣子：

    | 名字 | 差分 | 頭部縮圖 | 台詞 |

  · 一句一列，**照劇情順序排下來**。
  · **段落標題列**（進入廢城／A 骨龕…）：深灰底白粗字，四欄合併。
  · **演出列**（派生、時間差分、SE、進入戰鬥、感應動畫…）：名字與差分空著，
    文字寫在「台詞」欄，灰色斜體 —— 順序不動，它們就是劇情的一部分。
  · **只有差分沒有台詞的拍**：台詞欄寫「（無台詞，差分拍）」。
  · **主角空白**：名字「主角」、差分空、台詞「（主角空白）」，沒有縮圖。
  · **名字與差分分兩欄**（Ray 指定），差分欄寫 speakers.js 的 expr 鍵。

⚠⚠⚠ **Ray 改過的表才是正本**（2026-09-28 的教訓：他在 Excel 裡改完差分丟回來，
  我沒讀就從稿重新生成，把他的修改整份蓋掉，最後靠 Excel 的自動回復檔才救回來）。
  ⇒ 表一旦交出去，之後**只准走 `--refresh`**（讀他那一份、只換縮圖、文字一格不動），
    **不准再從稿重出一次去蓋它**。`--refresh` 蓋檔之前會先把原檔送進回收區（§5）。

⚠ 縮圖**不自己發明裁法**：沿用 `tools/si_xlsx.py` 的 `thumb()` ＝遊戲裡頭像
  （`speakers.faceStyle()`）看到的那一塊（鐵律 7）。

稿的格式（.txt，一行一拍）：

    #段落標題
    @演出說明（SE／派生／進入戰鬥…）
    諾|scare|這裡的天空……          ← 名字|差分|台詞；台詞空白＝差分拍
    主|-|（主角空白）

  名字可以寫簡稱（諾／蕾／安／索）、speakers.js 的顯示名、或 ART 的鑰匙。
  差分可以寫 expr 鍵、`base`，或直接寫檔名（`sorana_si_back.webp`）。
"""
import argparse, io, os, re, shutil, subprocess, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import si_xlsx
from si_xlsx import ROOT, load_speakers, name_table, frames, thumb, strip_v

TH = 88
HEAD = ['名字', '差分', '頭部縮圖', '台詞']
NOLINE = '（無台詞，差分拍）'
SHORT = {'諾': 'nouvelle', '蕾': 'renna', '安': 'anya', '索': 'sorana'}


def resolver():
    ART, SP = load_speakers()
    names = name_table(SP)                      # art → 顯示名
    back = {v: k for k, v in names.items()}     # 顯示名 → art
    F = {}
    for key, en, f in frames(ART):
        F[(key, 'base' if en == '(base)' else en)] = f
        F.setdefault((key, os.path.basename(strip_v(f.get('src', ''))).lower()), f)

    def art_of(who):
        return SHORT.get(who) or back.get(who) or (who if who in ART else None)

    def disp(key, who):
        return names.get(key) or who

    def frame(key, ex):
        if key is None or not ex:
            return None
        return F.get((key, ex)) or F.get((key, os.path.basename(ex).lower()))
    return art_of, disp, frame


def style_sheet(ws):
    from openpyxl.styles import Font, PatternFill
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
    for col, w in zip('ABCD', (10, 14, 14, 62)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'


def write(rows, path, art_of, disp, frame):
    """rows：[('#',標題) | ('@',說明) | ('主',台詞) | ('line',名字,差分,台詞)]"""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImg
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = os.path.splitext(os.path.basename(path))[0].split('_')[0][:31]
    ws.append(HEAD)
    style_sheet(ws)
    miss, keep = [], []
    for row in rows:
        r = ws.max_row + 1
        kind = row[0]
        if kind == '#':
            ws.cell(r, 1, row[1]).font = Font(bold=True, color='FFFFFF')
            for c in range(1, 5):
                ws.cell(r, c).fill = PatternFill('solid', fgColor='444444')
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            continue
        if kind == '@':
            ws.cell(r, 4, row[1]).font = Font(italic=True, color='888888')
            continue
        if kind == '主':
            ws.cell(r, 1, '主角')
            ws.cell(r, 4, row[1] or '（主角空白）')
            continue
        _, who, ex, tx = row
        key = art_of(who)
        ws.cell(r, 1, disp(key, who) if key else who)
        ws.cell(r, 2, ex)
        ws.cell(r, 4, tx or NOLINE)
        f = frame(key, ex)
        if not f:
            miss.append((who, ex))
            continue
        p = os.path.join(ROOT, strip_v(f['src']))
        if not os.path.exists(p):
            miss.append((who, ex + '（檔案不見）'))
            continue
        im = thumb(p, f, measured=not f.get('unmeasured'))
        im.thumbnail((TH, TH))
        buf = io.BytesIO()
        im.save(buf, 'PNG')
        keep.append(buf)
        ws.add_image(XImg(buf), 'C%d' % r)
        ws.row_dimensions[r].height = 70
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical='center', wrap_text=True)
    wb.save(path)
    return miss


def parse_txt(path):
    rows = []
    for ln in io.open(path, encoding='utf-8').read().splitlines():
        ln = ln.rstrip()
        if not ln:
            continue
        if ln[0] == '#':
            rows.append(('#', ln[1:].strip()))
        elif ln[0] == '@':
            rows.append(('@', ln[1:].strip()))
        else:
            who, ex, tx = (ln.split('|') + ['', ''])[:3]
            if who == '主':
                rows.append(('主', tx))
            else:
                rows.append(('line', who.strip(), ex.strip(), tx.strip()))
    return rows


def parse_xlsx(path):
    """讀 Ray 那一份：文字原樣帶回去，一格都不改。"""
    import openpyxl
    ws = openpyxl.load_workbook(path).active
    it = list(ws.iter_rows(values_only=True))
    if [str(x or '') for x in it[0][:4]] != HEAD:
        sys.exit('表頭不是標準格式 %r —— 不動它' % (it[0],))
    rows = []
    for a, b, _c, d in ((tuple(r) + (None,) * 4)[:4] for r in it[1:]):
        if a and not b and not d and a != '主角':
            rows.append(('#', a))
        elif not a and not b:
            if d:
                rows.append(('@', d))
        elif a == '主角':
            rows.append(('主', d))
        else:
            rows.append(('line', a, b or '', '' if d == NOLINE else (d or '')))
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('src', nargs='?')
    ap.add_argument('-o', '--out')
    ap.add_argument('--refresh')
    ap.add_argument('--append')
    a = ap.parse_args()
    art_of, disp, frame = resolver()
    if a.append and not a.src:
        sys.exit('用法：dialogue_xlsx.py 稿.txt --append 表.xlsx')
    if a.refresh or a.append:
        path = os.path.abspath(a.refresh or a.append)
        rows = parse_xlsx(path) + (parse_txt(a.src) if a.append else [])
        tmp = path + '.tmp.xlsx'
        miss = write(rows, tmp, art_of, disp, frame)
        # 蓋原檔之前先把 Ray 那一份送進回收區（§5：覆蓋一律走 recycle）
        subprocess.check_call([os.path.join(ROOT, 'tools', 'recycle.sh'), '-m',
                               ('dialogue_xlsx --append：接新稿前的原檔' if a.append
                                else 'dialogue_xlsx --refresh：重畫縮圖前的原檔'), path])
        shutil.move(tmp, path)
    else:
        if not a.src or not a.out:
            sys.exit('用法：dialogue_xlsx.py 稿.txt -o 表.xlsx ／ --refresh 表.xlsx')
        if os.path.exists(a.out):
            sys.exit('%s 已經存在 —— 交出去的表只准 --refresh，不准從稿重出去蓋它' % a.out)
        path = a.out
        miss = write(parse_txt(a.src), path, art_of, disp, frame)
    print('寫出', path)
    for m in miss:
        print('  ⚠ 找不到差分：', *m)


if __name__ == '__main__':
    main()
