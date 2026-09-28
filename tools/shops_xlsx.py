#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""商店卡 ⇄ Excel（ver -1831，Ray：「城鎮就算還沒接店主圖也要先有商店功能，目前先通用帝都版本，
   並製作 excel 商店卡之後我在表裡直接改了匯入」）

    python3 tools/shops_xlsx.py export           # script/shopcards.js → shops.xlsx（專案根目錄）
    python3 tools/shops_xlsx.py import [檔案]     # shops.xlsx → 重新產生 script/shopcards.js

⚠⚠⚠ **`import` 只有 Ray 明講的時候才跑**（同 enemies_xlsx.py 那條，ver -944）——
  那份 Excel 可能是他改到一半的草稿，沒說就抓等於拿半成品蓋掉線上的貨單。

⚠ **真相是 `script/shopcards.js`**：Excel 是編輯用的視圖。要改之前先 `export` 拿最新的。
⚠ shopcards.js 是**這支工具產生的檔**（整檔重寫，與 enemies.js 的「就地改值」不同）——
  所以各店的說明不寫成註解，而是卡上的「備註」欄（`note`），Excel 上看得到也改得到，重產不會丟。
⚠ 價格不在這裡：道具與武器的價格在各自的卡上（items.defs／weapons），表上的「參考價」只給看、匯入不讀。

兩張工作表：
  店    —— 一列一家店：key／城／店名／分頁／分頁名／只收／比較／射擊挑戰／挑戰名／折扣旗／折扣倍率／店主圖／備註
  貨單  —— 一列一項：店 key／品項 id／（品名，參考）／（參考價）／存貨（空白＝不限量）
"""
import json, os, sys
import _jsrun
import _utf8  # noqa: F401

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JS   = os.path.join(ROOT, 'script', 'shopcards.js')
XLSX = os.path.join(ROOT, 'shops.xlsx')

SHOP_COLS = [('key','key'), ('city','城'), ('title','店名'), ('tabs','分頁'), ('tabName','分頁名'),
             ('only','只收'), ('compare','比較'), ('challenge','射擊挑戰'), ('challengeLabel','挑戰名'),
             ('saleNeed','折扣旗'), ('saleMul','折扣倍率'), ('art','店主圖'), ('note','備註')]
STOCK_COLS = [('shop','店 key'), ('id','品項 id'), ('name','品名（參考）'), ('price','參考價'), ('n','存貨（空白＝不限量）')]


def load_cards():
    src = (f"import {{ SHOP_CARDS }} from '{_jsrun.file_url(JS)}';"
           f"print(JSON.stringify(SHOP_CARDS));")
    return _jsrun.dump(src, module=True)


def load_items():
    src = (f"import {{ GAME_CONFIG }} from '{_jsrun.file_url(os.path.join(ROOT,'config.js'))}';"
           f"const d=(GAME_CONFIG.items||{{}}).defs||{{}}, w=GAME_CONFIG.weapons||{{}}, o={{}};"
           f"for(const k in d) o[k]={{name:d[k].name, price:d[k].price}};"
           f"for(const k in w) o[k]={{name:w[k].name, price:(w[k].price!=null?w[k].price:w[k].value)}};"
           f"print(JSON.stringify(o));")
    return _jsrun.dump(src, module=True)


def js_val(v):
    return json.dumps(v, ensure_ascii=False)


def write_js(cards):
    """整檔重寫 script/shopcards.js。鍵的順序照 cards 的順序。"""
    out = ['/* ============================================================================',
           ' *  script/shopcards.js — 商店卡（唯一資料來源，ver -1831）',
           ' *  ---------------------------------------------------------------------------',
           ' *  ⚠⚠ **這一檔由 `tools/shops_xlsx.py import` 產生 —— 不要手改**（改了下次匯入會被蓋掉）。',
           ' *     要改：`python3 tools/shops_xlsx.py export` → 改 shops.xlsx → Ray 說匯入才 `import`。',
           ' *  ⚠ 各店的說明寫在卡上的 `note`（Excel 的「備註」欄），不寫成註解（整檔重產會洗掉註解）。',
           ' *  ⚠ 欄位的意義見 config.js 的 `shop` 那一段；節點用 `shop:\'<key>\'` 指到這裡的鍵。',
           ' *  ⚠ 貨單分開記帳：`script/shopstock.js` 的鑰匙就是這裡的 key（兩家店＝兩本帳）。',
           ' *  ⚠ 純資料檔，不 import 任何東西（不會與 config 成環）。',
           ' * ========================================================================== */',
           'export const SHOP_CARDS = {']
    for k, c in cards.items():
        out.append('  %s: {' % k)
        for f in ['city', 'title', 'art', 'tabs', 'tabName', 'only', 'compare', 'challenge', 'challengeLabel', 'sale', 'note']:
            if c.get(f) in (None, '', [], {}):
                continue
            out.append('    %s: %s,' % (f, js_val(c[f])))
        stock = c.get('stock') or []
        out.append('    stock: [')
        for s in stock:
            out.append('      %s,' % js_val(s))
        out.append('    ],')
        out.append('  },')
    out.append('};')
    tmp = JS + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        f.write('\n'.join(out) + '\n')
    os.replace(tmp, JS)


def export():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    cards = load_cards()
    items = load_items()
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = '店'
    ws.append([h for _, h in SHOP_COLS])
    for k, c in cards.items():
        sale = c.get('sale') or {}
        row = {'key': k, 'city': c.get('city', ''), 'title': c.get('title', ''),
               'tabs': ','.join(c.get('tabs') or []),
               'tabName': json.dumps(c['tabName'], ensure_ascii=False) if c.get('tabName') else '',
               'only': c.get('only', ''), 'compare': 1 if c.get('compare') else '',
               'challenge': c.get('challenge', ''), 'challengeLabel': c.get('challengeLabel', ''),
               'saleNeed': sale.get('need', ''), 'saleMul': sale.get('mul', ''),
               'art': c.get('art', ''), 'note': c.get('note', '')}
        ws.append([row[f] for f, _ in SHOP_COLS])
    ws2 = wb.create_sheet('貨單')
    ws2.append([h for _, h in STOCK_COLS])
    for k, c in cards.items():
        for s in c.get('stock') or []:
            sid, n = (s, '') if isinstance(s, str) else (s.get('id'), s.get('n', ''))
            it = items.get(sid) or {}
            ws2.append([k, sid, it.get('name', '（查不到）'), it.get('price', ''), n])
    for w in (ws, ws2):
        for cell in w[1]:
            cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='EEDDAA')
        w.freeze_panes = 'A2'
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(10, min(48, max(len(str(c.value or '')) for c in col) * 1.6))
    # 參考欄標灰：匯入不讀
    for cell in ws2['C'] + ws2['D']:
        cell.font = Font(color='888888')
    wb.save(XLSX)
    print('匯出 %s（%d 家店、%d 項貨）' % (XLSX, len(cards), sum(len(c.get('stock') or []) for c in cards.values())))


def cell(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v.strip() if isinstance(v, str) else v


def do_import(path=None):
    import openpyxl
    wb = openpyxl.load_workbook(path or XLSX, data_only=True)
    ws, ws2 = wb['店'], wb['貨單']
    hdr = [cell(c.value) for c in ws[1]]
    col = {f: hdr.index(h) for f, h in SHOP_COLS if h in hdr}
    cards, errs = {}, []
    for r in ws.iter_rows(min_row=2, values_only=True):
        v = {f: cell(r[i]) if i < len(r) else '' for f, i in col.items()}
        k = v.get('key')
        if not k:
            continue
        c = {'city': v.get('city') or '', 'title': v.get('title') or '', 'art': v.get('art') or '',
             'tabs': [t.strip() for t in str(v.get('tabs') or '').split(',') if t.strip()]}
        if v.get('tabName'):
            try:
                c['tabName'] = json.loads(v['tabName'])
            except Exception:
                errs.append('%s 的「分頁名」不是合法的 JSON：%s' % (k, v['tabName']))
        for f in ('only', 'challenge', 'challengeLabel', 'note'):
            if v.get(f) not in ('', None):
                c[f] = v[f]
        if v.get('compare') not in ('', None, 0, '0'):
            c['compare'] = True
        if v.get('saleNeed'):
            c['sale'] = {'need': v['saleNeed'], 'mul': float(v.get('saleMul') or 1)}
        c['stock'] = []
        cards[k] = c
    hdr2 = [cell(c.value) for c in ws2[1]]
    col2 = {f: hdr2.index(h) for f, h in STOCK_COLS if h in hdr2}
    for r in ws2.iter_rows(min_row=2, values_only=True):
        shop = cell(r[col2['shop']]); sid = cell(r[col2['id']])
        if not shop or not sid:
            continue
        if shop not in cards:
            errs.append('貨單裡的店 %s 不在「店」那一張' % shop); continue
        n = cell(r[col2['n']]) if 'n' in col2 else ''
        if n in ('', None):
            cards[shop]['stock'].append(sid)
        else:
            try:
                cards[shop]['stock'].append({'id': sid, 'n': int(n)})
            except Exception:
                errs.append('%s / %s 的存貨不是整數：%s' % (shop, sid, n))
    if errs:
        print('⚠ 有問題，沒寫入：'); [print('  ' + e) for e in errs]; sys.exit(1)
    write_js(cards)
    print('已匯入 → %s（%d 家店）' % (JS, len(cards)))


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else ''
    if cmd == 'export':
        export()
    elif cmd == 'import':
        do_import(sys.argv[2] if len(sys.argv) > 2 else None)
    elif cmd == '_gen':                      # 內部：從 JSON 產生（第一次搬家用）
        write_js(json.load(open(sys.argv[2], encoding='utf-8')))
    else:
        print(__doc__)
