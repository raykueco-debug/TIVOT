#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""九星技能表 → Excel（ver -1011，Ray：「把技能表出個 excel 給我，有數值的地方獨立欄」）

    python3 tools/girlstars_xlsx.py          # config.js → girlstars.xlsx（專案根目錄）

⚠⚠ **真相是 `config.js` 的 `GAME_CONFIG.girls.levels`**（鐵律 1/7）——
  這份 Excel 是**看與討論用的視圖**，不是第二份資料。這支工具目前**只有匯出**：
  要改數值請改 config.js（或跟我說要改哪一格）。
  ⚠ 日後若要做 `import`，照 `tools/enemies_xlsx.py` 那一套「就地改值、不重產檔案」
    寫 —— 星表那邊的 ⚠ 註解（哪個數字是 Ray 指定的、為什麼這樣調）是最貴的東西。

⚠ 讀 config.js 用 macOS 內建的 `jsc` 跑一次真的 import（同 script_lint.py 的作法）——
  自己寫 JS 解析器一定會在某個引號或巢狀上翻車。
"""
import json, os, re, subprocess, sys, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JSC  = '/System/Library/Frameworks/JavaScriptCore.framework/Versions/A/Helpers/jsc'
XLSX = os.path.join(ROOT, 'girlstars.xlsx')
# ⚠ speakers.js 要一起載：config.js 的 `tutPortraits` 那一段會讀 `ART`。
SRC  = ('script/speakers.js', 'script/enemies.js', 'config.js')

WHO_CN  = {'nouvelle': '諾薇兒', 'anya': '安雅', 'sorana': '索菈娜'}
SKILL_CN = {'install': '覺醒技', 'active': '主動技', 'passive': '被動技'}

# ── 每個數值鍵是什麼意思（Excel 第 2 列的小字）─────────────────────────
#   ⚠ 只寫「這一格在講什麼」，**不要抄一份數字進來**（鐵律 7：值只有 config 一份）。
GLOSS = {
  # 諾薇兒
  'saintComboMul':    '聖徒化連擊疊傷斜率的**增量**（0.5＝×1.5）',
  'saintHint':        '聖徒化期間全程指引下一格',
  'saintReload':      '無傷擊殺回填聖徒化',
  'guardHealPct':     '獄門天鎖免傷期間每發回復（佔最大體力）',
  'guardHealCounter': '反擊也算一發（回血窗）',
  'lifeReturnSec':    '魂之歸所延續聖徒化連擊增傷的秒數',
  'lifeReturnCombo':  '魂之歸所延續連擊增傷',
  'lifeReturnHint':   '魂之歸所期間指引下一格',
  'saintStartHp1':    '聖徒化發動時體力降至 1',
  'guardPerEnemy':    '獄門天鎖每一場戰鬥都可發動',
  'saintNoHitAdvance':'受敵攻擊不推進倒數槽',
  'saintDrainPct':    '聖徒化期間每發延長倒數（佔最大體力）',
  # 安雅
  'niDmgMul':         '夢魘化期間攻擊力**增量**（0.2＝×1.2）',
  'burstBuffSec':     '夢境破碎後的反擊增益秒數',
  'burstAtk':         '該增益的攻擊力帶位（0黃 1橘 2紅，可累加）',
  'burstHint':        '該增益期間指引下一格',
  'niReload':         '連續三次完美反擊回填夢魘化',
  'lucidSec':         '明晰之夢秒數的**增量**（可累加）',
  'counterAtk':       '明晰之夢期間的攻擊力帶位（0黃 1橘 2紅，可累加）',
  'niFullStart':      '夢魘化發動時體力先回滿',
  'burstLastCell':    '最後一格夢境破碎的追加傷害（佔敵最大體力）',
  'niCounterPauseSec':'每次反擊讓夢魘化抽血暫停的秒數',
  # 索菈娜
  'brDmgMul':         '彈雨傾洩攻擊力**增量**（0.2＝×1.2）',
  'roarReloadActive': '獵手的戰吼發動時回填主動技',
  'activeEnergyBuffSec':'主動技後破防值累積 200% 的秒數',
  'coopAtk':          '共鬥自動反擊的攻擊力帶位（0黃 1橘 2紅，可累加）',
  'roarStreakCut':    '戰吼所需連續完美清盤數的**減量**',
  'activeReloadCoop': '主動技發動同時回填共鬥',
  'coopSec':          '共鬥最大持續時間的**增量**（秒）',
  'coopEnergyTime':   '共鬥期間可累積破防值（＝延長時間）',
}

def load():
    if not os.path.exists(JSC):
        print('找不到 jsc（%s）——這支工具依賴 macOS 內建的 JavaScriptCore。' % JSC); sys.exit(2)
    def strip(src):
        src = re.sub(r'^\s*import[^;]*;', '', src, flags=re.M)
        src = re.sub(r'^\s*export\s+(?=(const|let|var|function|class|async))', '', src, flags=re.M)
        src = re.sub(r'^\s*export\s*\{[^}]*\};?', '', src, flags=re.M)
        return src
    parts = [strip(open(os.path.join(ROOT, f), encoding='utf-8').read()) for f in SRC]
    parts.append('print(JSON.stringify({v:VERSION, girls:GAME_CONFIG.girls, partners:GAME_CONFIG.partners}));')
    t = tempfile.NamedTemporaryFile('w', suffix='.js', delete=False, encoding='utf-8')
    t.write('\n'.join(parts)); t.close()
    r = subprocess.run([JSC, t.name], capture_output=True, text=True); os.unlink(t.name)
    if r.returncode or not r.stdout.strip():
        print('讀不到 config：\n' + (r.stderr or r.stdout)); sys.exit(2)
    return json.loads(r.stdout)

def main():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter as CL

    d = load(); girls = d['girls']; parts = d['partners']
    order = [w for w in ('nouvelle', 'anya', 'sorana') if w in girls.get('levels', {})]

    # 數值欄：依「角色 → Lv」第一次出現的順序排（同一角色的欄會排在一起）
    keys = []
    for who in order:
        for st in girls['levels'][who]:
            for k in st:
                if k in ('star', 'name', 'skill', 'desc', 'en'): continue
                if k not in keys: keys.append(k)

    base = ['角色', 'Lv', '星（西文）', '星名', '歸屬', '效果說明']
    names = base + keys

    wb = Workbook(); ws = wb.active; ws.title = '九星'
    thin = Side(style='thin', color='C8B99A'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    grp = PatternFill('solid', fgColor='E8D9BC'); hdr = PatternFill('solid', fgColor='FFF2E0')
    numfill = PatternFill('solid', fgColor='EAF2FA')

    # 第 1 列：群組
    row1 = ['基本'] * len(base) + ['數值（每個效果鍵一欄）'] * len(keys)
    # 第 2 列：這一格在講什麼
    row2 = ['', '1~9', '星座上的星名', '中文星名', '覺醒／主動／被動', '技能表上印給玩家看的那一句'] + \
           [GLOSS.get(k, '') for k in keys]
    ws.append(row1); ws.append(row2); ws.append(names)

    for who in order:
        for i, st in enumerate(girls['levels'][who], 1):
            r = [WHO_CN.get(who, who), i, st.get('star', ''), st.get('name', ''),
                 SKILL_CN.get(st.get('skill', ''), st.get('skill', '')),
                 re.sub(r'<br>', ' / ', st.get('desc', ''))]
            for k in keys:
                v = st.get(k, None)
                r.append('' if v is None else v)
            ws.append(r)

    for c in ws[1]: c.font = Font(bold=True); c.fill = grp
    for c in ws[2]: c.font = Font(size=9, color='808080'); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in ws[3]: c.font = Font(bold=True); c.fill = hdr
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(names)):
        for c in row:
            c.border = box
            if c.row >= 4 and c.column > len(base):
                c.fill = numfill; c.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'G4'
    for i, n in enumerate(names, 1):
        w = max([len(str(n))] + [len(str(ws.cell(r, i).value or '')) for r in range(4, ws.max_row + 1)])
        ws.column_dimensions[CL(i)].width = min(max(w * 1.3 + 2, 8), 60)
    ws.column_dimensions[CL(names.index('效果說明') + 1)].width = 58
    ws.row_dimensions[2].height = 42

    # ── 第二張：基礎技能（點星之前的初始文案）──────────────────────
    ws2 = wb.create_sheet('基礎技能')
    ws2.append(['角色', '欄位', '名稱', '西文', '說明'])
    for c in ws2[1]: c.font = Font(bold=True); c.fill = hdr; c.border = box
    for who in order:
        p = parts.get(who, {})
        for f, label in (('install', '覺醒技'), ('active', '主動技'), ('passive', '被動技')):
            o = p.get(f)
            if not o: continue
            ws2.append([WHO_CN.get(who, who), label, o.get('name', ''), o.get('en', ''),
                        re.sub(r'<br>', ' / ', o.get('desc', ''))])
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, max_col=5):
        for c in row: c.border = box; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate([10, 8, 14, 20, 90], 1): ws2.column_dimensions[CL(i)].width = w

    wb.save(XLSX)
    print('已寫出 %s（%s）\n  九星 %d 列、數值欄 %d 欄；基礎技能 %d 列'
          % (os.path.relpath(XLSX, ROOT), d.get('v', ''), ws.max_row - 3, len(keys), ws2.max_row - 1))

if __name__ == '__main__':
    main()
