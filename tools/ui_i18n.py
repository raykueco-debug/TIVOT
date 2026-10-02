#!/usr/bin/env python3
"""tools/ui_i18n.py —— 介面字多語（ver -1909）

劇本（對白／地名）走 tools/script_i18n.py；**其餘所有玩家看得到的中文**走這一支：
引擎程式裡寫死的字、資料卡（敵人／武器／道具／商店）、index.html 與飛行頁的介面。

    python3 tools/ui_i18n.py extract          # 盤點 → i18n/ui/strings.json ＋ i18n/ui/INVENTORY.md（統計）
    python3 tools/ui_i18n.py wrap             # 把程式碼裡的中文字面值包成 i18nT('…')（只做一次；冪等）
    python3 tools/ui_i18n.py check es         # 驗 i18n/ui/es.json
    python3 tools/ui_i18n.py xlsx es          # → reference/ui_es.xlsx（給人審）
    python3 tools/ui_i18n.py js es            # → i18n/ui/es.js（遊戲讀的那一份）

⚠ 新增語種：複製 i18n/ui/es.json 的鑰匙（＝strings.json 的 zh）翻成新語言 → `js <lang>` →
  在 i18n/scriptTr.js 的 AVAILABLE 加上它。**盤點與包裝不必重做**（那是這一次的成果）。

每一處中文的「處理方式」（mode）：
  wrap    程式字面值 → 執行時 i18nT('…') 查表（劇本表＋介面表兩張都查）
  data    資料檔（config.js 與它 import 的卡）→ config.js 檔尾 trTree(GAME_CONFIG) 就地換
  html    index.html／飛行頁的靜態標記 → 開機時 trDom() 掃文字節點與屬性
  script  已經在劇本譯文表裡（tools/script_i18n.py）—— 不重複
  skip-*  不譯：console／物件鍵／case／比較運算元／含 ${} 的樣板字串／著色器／識別字
"""
import re, sys, json, os, glob
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from script_i18n import tokens, unesc, CJK as _CJK, ROOT
# 介面這邊連「只有全形標點」的字串也要（`'」'`、`'（'`、`'：'` 會跟中文黏在一起顯示）；純空白不算
class _W:
    _p = re.compile(r'[\u3001-\u303f\uff01-\uff5e]')
    @staticmethod
    def search(s): return _CJK.search(s) or _W._p.search(s)
CJK = _W

OUT = os.path.join(ROOT, 'i18n', 'ui')

CODE = sorted(glob.glob(os.path.join(ROOT, 'modules', '*.js'))) + [os.path.join(ROOT, p) for p in (
    'main.js', 'config.js', 'script/clock.js', 'script/progress.js', 'script/inventory.js',
    'script/loadout.js', 'script/shopstock.js')]
# 純資料卡（不 import 任何東西、shopcards 還是工具產生的）→ 不改檔，由 config.js 檔尾 trTree(GAME_CONFIG) 換
DATA = [os.path.join(ROOT, p) for p in (
    'script/enemies.js', 'script/weapons.js', 'script/shopcards.js')]
# 拿來當識別字、但也會顯示的字（顯示端寫 T(變數)）—— 強制進表
FORCE = ['重機槍', '霰彈槍', '萊福槍'] + ['%d月' % m for m in range(1, 13)]   # 月名：clock 的日期樣板 {M}
HTML = [os.path.join(ROOT, p) for p in ('index.html', 'flight/index.html')]

# 拿中文當識別字、不可以換的欄位（同 i18n/scriptTr.js 的 SKIP）
ID_KEYS = {'time', 'art', 'expr', 'img', 'bg', 'cg', 'se', 'bgm', 'flag', 'flags', 'need', 'until', 'cat', 'except'}

def rel(p): return os.path.relpath(p, ROOT).replace(os.sep, '/')

def classify(src, pos, raw, q, end):
    back = src[max(0, pos-400):pos]
    # 判鍵／比較前先把註解拿掉（`'霰彈槍':{` 前面常常隔著一段 /* … */）
    back = re.sub(r'/\*.*?\*/', ' ', back, flags=re.S)
    back = re.sub(r'//[^\n]*', ' ', back)
    after = src[end:end+40]
    if re.search(r'console\.\w+\([^;]*$', back) or re.search(r'\bthrow\b[^;]*$', back):
        return 'skip-console'
    if raw.lstrip().startswith('#version') or 'precision highp' in raw:
        return 'skip-shader'
    if q == '`' and '${' in raw:
        return 'skip-template'
    if re.match(r'\s*:', after) and re.search(r'[{,]\s*$', back):   # 物件鍵（註解已先拿掉；三元運算的 `? 'a' : b` 不算）
        return 'skip-key'
    if re.search(r'\bcase\s*$', back):
        return 'skip-case'
    if re.search(r'(===|!==|==|!=)\s*$', back) or re.match(r'\s*(===|!==|==|!=)', after):
        return 'skip-compare'
    if re.search(r'\b(indexOf|includes|startsWith|endsWith|split|replace|replaceAll|has|get|set)\(\s*$', back):
        return 'skip-compare'
    km = re.search(r'([A-Za-z_$][\w$]*)\s*:\s*\[?\s*$', back[-60:])
    if km and km.group(1) in ID_KEYS:
        return 'skip-id'
    return None

def scan_js(path, src, base_line, mode, items, script_tab):
    for pos, raw, q in tokens(src):
        if not CJK.search(raw): continue
        end = pos + len(raw) + 2
        m = classify(src, pos, raw, q, end) or mode
        text = unesc(raw) if q != '`' else raw
        if m in ('wrap', 'data') and text in script_tab: m = 'script'
        line = base_line + src.count('\n', 0, pos)
        it = items.setdefault(text, {'zh': text, 'occ': []})
        it['occ'].append({'f': rel(path), 'l': line, 'm': m, 'q': q})

def blank_comments(s):
    return re.sub(r'<!--.*?-->', lambda m: re.sub(r'[^\n]', ' ', m.group(0)), s, flags=re.S)

SCRIPT_RE = r'<script(?![^>]*\bsrc=)(?![^>]*importmap)[^>]*>(.*?)</script>'

def scan_html(path, items, script_tab):
    s = open(path, encoding='utf-8').read()
    # 行內 <script>（非 src、非 importmap）→ 程式；⚠ 先遮掉 HTML 註解（註解裡會寫到 `<script>` 字樣）
    for m in re.finditer(SCRIPT_RE, blank_comments(s), re.S):
        scan_js(path, m.group(1), s.count('\n', 0, m.start(1)) + 1, 'wrap', items, script_tab)
    body = s
    for m in re.finditer(r'<script.*?</script>|<style.*?</style>|<!--.*?-->', s, re.S):
        body = body[:m.start()] + ' ' * (m.end()-m.start()) + body[m.end():]
    for m in re.finditer(r'>([^<]+)<', body):
        t = m.group(1).strip()
        if CJK.search(t):
            it = items.setdefault(t, {'zh': t, 'occ': []})
            it['occ'].append({'f': rel(path), 'l': s.count('\n', 0, m.start()) + 1, 'm': 'script' if t in script_tab else 'html'})
    for m in re.finditer(r'\b(placeholder|title|aria-label|alt)="([^"]*)"', body):
        t = m.group(2).strip()
        if CJK.search(t):
            it = items.setdefault(t, {'zh': t, 'occ': []})
            it['occ'].append({'f': rel(path), 'l': s.count('\n', 0, m.start()) + 1, 'm': 'html'})

def extract():
    script_tab = set(json.load(open(os.path.join(ROOT, 'i18n', 'script', 'es.json'), encoding='utf-8')))
    items = {}
    for p in CODE: scan_js(p, open(p, encoding='utf-8').read(), 1, 'wrap', items, script_tab)
    for p in DATA: scan_js(p, open(p, encoding='utf-8').read(), 1, 'data', items, script_tab)
    for p in HTML: scan_html(p, items, script_tab)
    for t in FORCE:
        items.setdefault(t, {'zh': t, 'occ': []})['occ'].append({'f': '(FORCE)', 'l': 0, 'm': 'data'})
    lst = list(items.values())
    for it in lst:
        ms = {o['m'] for o in it['occ']}
        it['need'] = bool(ms & {'wrap', 'data', 'html'})
    os.makedirs(OUT, exist_ok=True)
    json.dump(lst, open(os.path.join(OUT, 'strings.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    write_inventory(lst)
    need = [x for x in lst if x['need']]
    print(f'{len(lst)} 種字串；要翻 {len(need)} 種（{sum(len(x["zh"]) for x in need)} 字）→ i18n/ui/strings.json、INVENTORY.md')

def write_inventory(lst):
    from collections import Counter, defaultdict
    by = defaultdict(Counter); ch = defaultdict(Counter)
    for it in lst:
        for o in it['occ']:
            by[o['f']][o['m']] += 1; ch[o['f']][o['m']] += len(it['zh'])
    modes = ['wrap', 'data', 'html', 'script', 'skip-console', 'skip-key', 'skip-case',
             'skip-compare', 'skip-template', 'skip-shader', 'skip-id']
    L = ['# 介面字盤點（tools/ui_i18n.py extract 產生，不要手改）', '',
         '> 這是**語種無關**的盤點：哪裡有中文、各用什麼方式換。新增語種只要翻',
         '> `i18n/ui/strings.json` 裡 `need:true` 的那些 `zh`，不必重新盤點。', '',
         '| 檔案 | ' + ' | '.join(modes) + ' |', '|---|' + '---|' * len(modes)]
    tot = Counter()
    for f in sorted(by):
        L.append('| `%s` | ' % f + ' | '.join(str(by[f][m] or '') for m in modes) + ' |')
        tot.update(by[f])
    L.append('| **合計（處）** | ' + ' | '.join('**%d**' % tot[m] for m in modes) + ' |')
    need = [x for x in lst if x['need']]
    L += ['', f'- 不重複字串：{len(lst)} 種；要翻：**{len(need)} 種／{sum(len(x["zh"]) for x in need)} 字**',
          '- `skip-template`／`skip-compare` 的那幾處是**人工檢查清單**：含 `${}` 的樣板字串查不了表、',
          '  比較運算元換了會壞邏輯 —— 玩家看得到的要改寫成 `T()` 能查的形狀。', '',
          '## 人工檢查清單（skip-template／skip-compare）', '']
    for it in lst:
        for o in it['occ']:
            if o['m'] in ('skip-template', 'skip-compare'):
                L.append(f"- `{o['f']}:{o['l']}` {o['m']}：{it['zh'][:60]!r}")
    open(os.path.join(OUT, 'INVENTORY.md'), 'w', encoding='utf-8').write('\n'.join(L) + '\n')

# ── 包裝：程式碼裡的字面值 → T('…') ─────────────────────────────────────────
IMPORT_MOD = "import { i18nT } from '%s';   // 介面字譯文（ver -1909；中文時原樣回傳）\n"

def wrap():
    data = json.load(open(os.path.join(OUT, 'strings.json'), encoding='utf-8'))
    want = {}   # (file, line) → set of zh to wrap
    for it in data:
        for o in it['occ']:
            # 'script' ＝ 字已在劇本譯文表裡，但若是**程式字面值**照樣要包（T 兩張表都查）
            code_file = o['f'].endswith('.html') or os.path.join(ROOT, o['f']) in CODE
            if o['m'] == 'wrap' or (o['m'] == 'script' and o.get('q') and code_file): want.setdefault(o['f'], set()).add(it['zh'])
    n_total = 0
    for f, texts in sorted(want.items()):
        path = os.path.join(ROOT, f); s = open(path, encoding='utf-8').read()
        if f.endswith('.html'):
            segs = [(m.start(1), m.end(1)) for m in re.finditer(SCRIPT_RE, blank_comments(s), re.S)]
        else:
            segs = [(0, len(s))]
        out = []; last = 0; n = 0
        for a, b in segs:
            src = s[a:b]
            for pos, raw, q in tokens(src):
                if not CJK.search(raw): continue
                end = pos + len(raw) + 2
                if classify(src, pos, raw, q, end): continue
                text = unesc(raw) if q != '`' else raw
                if text not in texts: continue
                if re.search(r'\bi18nT\(\s*$', src[max(0, pos-8):pos]): continue      # 已包過
                g0 = a + pos; g1 = a + end
                out.append(s[last:g0]); out.append('i18nT(' + s[g0:g1] + ')'); last = g1; n += 1
        out.append(s[last:]); s2 = ''.join(out)
        if n and not f.endswith('.html') and not re.search(r'import\s*\{[^}]*\bi18nT\b', s2):
            d = os.path.dirname(f); target = os.path.relpath('i18n/scriptTr.js', d or '.').replace(os.sep, '/')
            if not target.startswith('.'): target = './' + target
            m = re.search(r'^import .*$', s2, re.M)
            ins = IMPORT_MOD % target
            s2 = (s2[:m.start()] + ins + s2[m.start():]) if m else ins + s2
        if s2 != s:
            open(path, 'w', encoding='utf-8').write(s2); n_total += n
            print(f'{f}: 包了 {n} 處')
    print('合計', n_total)

# ── 譯文表 ─────────────────────────────────────────────────────────────────
def load(lang):
    src = [x for x in json.load(open(os.path.join(OUT, 'strings.json'), encoding='utf-8')) if x['need']]
    p = os.path.join(OUT, f'{lang}.json')
    tr = json.load(open(p, encoding='utf-8')) if os.path.exists(p) else {}
    return src, tr

def check(lang):
    src, tr = load(lang)
    zh = {x['zh'] for x in src}
    miss = [x['zh'] for x in src if x['zh'] not in tr]
    orphan = [k for k in tr if k not in zh]
    ph = re.compile(r'\{[A-Za-z0-9_+]+\}|<[^>]+>|\$\{[^}]*\}')
    # 允許的差異：給人讀的屬性值（aria-label／placeholder／title／alt）本來就要翻；日期樣板的 {m} 可換成月名 {M}
    norm = lambda t: sorted(re.sub(r'\b(aria-label|placeholder|title|alt)="[^"]*"', r'\1=""', x).replace('{M}', '{m}') for x in ph.findall(t))
    bad = [k for k in tr if k in zh and norm(k) != norm(tr[k]) and set(norm(k)) != set(norm(tr[k]))]
    KEEP_CJK = {'日本語', '中文', '月'}   # 語言鈕上「下一個語言的自稱」／月名樣板的中間鍵（i18nT(m+'月')）
    cjk = [k for k in tr if k in zh and k not in KEEP_CJK and re.search(r'[\u4e00-\u9fff]', tr[k])]
    print(f'{lang}: {len(src)-len(miss)}/{len(src)} 已譯｜缺 {len(miss)}｜孤兒 {len(orphan)}｜標籤/佔位符不符 {len(bad)}｜殘留中文 {len(cjk)}')
    for k in bad[:15]: print('  不符', repr(k), '→', repr(tr[k]))
    for k in cjk[:15]: print('  殘留', repr(k), '→', repr(tr[k]))
    return 1 if (miss or bad) else 0

def xlsx(lang):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    src, tr = load(lang)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = lang
    ws.append(['#', '處理方式', '位置', '中文', lang])
    for i, x in enumerate(src, 1):
        o = x['occ'][0]
        ws.append([i, ','.join(sorted({p['m'] for p in x['occ']})), f"{o['f']}:{o['l']}" + (f" (+{len(x['occ'])-1})" if len(x['occ']) > 1 else ''),
                   x['zh'], tr.get(x['zh'], '')])
    for c, w in zip('ABCDE', (6, 10, 30, 60, 70)): ws.column_dimensions[c].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[3:]: c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in ws[1]: c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    p = os.path.join(ROOT, 'reference', f'ui_{lang}.xlsx'); wb.save(p); print('→', rel(p))

def js(lang):
    src, tr = load(lang)
    out = {x['zh']: tr[x['zh']] for x in src if x['zh'] in tr}
    p = os.path.join(OUT, f'{lang}.js')
    with open(p, 'w', encoding='utf-8') as f:
        f.write('/* 由 tools/ui_i18n.py js %s 從 %s.json 產生 —— 不要手改，改 json 再重跑。 */\n' % (lang, lang))
        f.write('export default ' + json.dumps(out, ensure_ascii=False, separators=(',', ':')) + ';\n')
    # 飛行頁是非模組、要同步載 → 另出一份 classic（掛到 window.__TIVOT_TR，兩張表合併）
    pc = p[:-3] + '.classic.js'
    with open(pc, 'w', encoding='utf-8') as f:
        f.write('/* 由 tools/%s js %s 產生（飛行頁用的 classic 版）—— 不要手改。 */\n' % (os.path.basename(__file__), lang))
        f.write('window.__TIVOT_TR=Object.assign(window.__TIVOT_TR||{},' + json.dumps(out, ensure_ascii=False, separators=(',', ':')) + ');\n')
    print(f'→ {rel(p)}（{len(out)} 句，{os.path.getsize(p)//1024} KB）')

if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'extract'
    {'extract': extract, 'wrap': wrap}.get(cmd, lambda: None)()
    if cmd == 'check': sys.exit(check(sys.argv[2]))
    if cmd == 'xlsx': xlsx(sys.argv[2])
    if cmd == 'js': js(sys.argv[2])
